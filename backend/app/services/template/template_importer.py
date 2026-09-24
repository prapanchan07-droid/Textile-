"""Imports a filled data template into the database so every dashboard page can read it."""
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.ingestion import (
    ReportRecord, ProductionRecord, DowntimeRecord, ManpowerRecord,
)
from app.models.template_data import (
    MachineDailyRecord, QualityParameterRecord, BusinessDailyRecord,
    StockRecord, ActionItemRecord,
)
from app.services.ingestion.normalizer import Normalizer
from app.services.template.template_spec import (
    SHEETS, DATA_SHEET_NAMES, TEMPLATE_MARKER_SHEET, TEMPLATE_REPORT_TYPE,
    MACHINE_SHEET, MANPOWER_SHEET, QUALITY_SHEET, BUSINESS_SHEET, STOCK_SHEET, ACTIONS_SHEET,
)

MAX_WARNINGS = 25
# Tables whose rows carry report_date + source_report_id and are replaced per date on re-upload
_REPLACEABLE_MODELS = [
    MachineDailyRecord, ProductionRecord, DowntimeRecord, ManpowerRecord,
    QualityParameterRecord, BusinessDailyRecord, StockRecord, ActionItemRecord,
]


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or (isinstance(v, str) and not v.strip())


def _text(v: Any) -> Optional[str]:
    if _blank(v):
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _num(v: Any, default: float = 0.0) -> float:
    if _blank(v):
        return default
    return Normalizer.parse_number(v, default)


def _opt_num(v: Any) -> Optional[float]:
    return None if _blank(v) else Normalizer.parse_number(v, 0.0)


def _date(v: Any) -> Optional[str]:
    if _blank(v):
        return None
    try:
        return pd.to_datetime(v).strftime("%Y-%m-%d")
    except Exception:
        return None


class TemplateImporter:
    @staticmethod
    def purge_report(db: Session, report_id: str) -> None:
        """Delete every dashboard row that came from one uploaded report."""
        for model in _REPLACEABLE_MODELS:
            db.query(model).filter(model.source_report_id == report_id).delete(synchronize_session=False)

    @staticmethod
    def is_template(file_path: str) -> bool:
        try:
            wb = load_workbook(file_path, read_only=True)
            names = set(wb.sheetnames)
            wb.close()
        except Exception:
            return False
        return TEMPLATE_MARKER_SHEET in names and any(s in names for s in DATA_SHEET_NAMES)

    @staticmethod
    def _read_sheet(file_path: str, sheet: str) -> Tuple[List[Tuple[int, Dict[str, Any]]], List[str]]:
        """Returns [(excel_row_number, {header: value})] and warnings. Raises on missing required headers."""
        df = pd.read_excel(file_path, sheet_name=sheet)
        df.columns = [str(c).strip() for c in df.columns]
        wanted = {c.header.lower(): c for c in SHEETS[sheet]["columns"]}
        renamed = {c: wanted[c.lower()].header for c in df.columns if c.lower() in wanted}
        df = df.rename(columns=renamed)

        missing = [c.header for c in SHEETS[sheet]["columns"] if c.required and c.header not in df.columns]
        if missing and not df.dropna(how="all").empty:
            raise ValueError(f"Sheet '{sheet}' is missing required column(s): {', '.join(missing)}.")

        rows = []
        for idx, rec in enumerate(df.to_dict("records")):
            if all(_blank(v) for v in rec.values()):
                continue
            rows.append((idx + 2, rec))
        return rows, []

    @staticmethod
    def import_file(db: Session, file_id: str, file_hash: str, filename: str, file_path: str) -> Dict[str, Any]:
        warnings: List[str] = []

        def warn(msg: str):
            if len(warnings) < MAX_WARNINGS:
                warnings.append(msg)

        wb = load_workbook(file_path, read_only=True)
        sheet_names = set(wb.sheetnames)
        wb.close()
        parsed: Dict[str, List[Tuple[int, Dict[str, Any]]]] = {}
        for sheet in DATA_SHEET_NAMES:
            if sheet in sheet_names:
                parsed[sheet] = TemplateImporter._read_sheet(file_path, sheet)[0]

        new_rows: List[Any] = []
        counts: Dict[str, int] = {}
        dates: set = set()

        def keep_date(sheet: str, row_no: int, rec: Dict[str, Any]) -> Optional[str]:
            d = _date(rec.get("Date"))
            if d is None:
                warn(f"{sheet} row {row_no}: missing or invalid Date - row skipped.")
                return None
            dates.add(d)
            return d

        for row_no, rec in parsed.get(MACHINE_SHEET, []):
            d = keep_date(MACHINE_SHEET, row_no, rec)
            mid, mtype = _text(rec.get("Machine ID")), _text(rec.get("Machine Type"))
            if d is None:
                continue
            if not mid or not mtype:
                warn(f"{MACHINE_SHEET} row {row_no}: Machine ID and Machine Type are required - row skipped.")
                continue
            target, actual = _num(rec.get("Target Kg")), _num(rec.get("Actual Kg"))
            shift = Normalizer.parse_shift(_text(rec.get("Shift")))
            unit = _text(rec.get("Unit"))
            section = _text(rec.get("Section")) or mtype
            dt_min = _num(rec.get("Downtime Min"))
            reason = _text(rec.get("Downtime Reason"))
            planned = (_text(rec.get("Planned Downtime")) or "N").upper().startswith("Y")
            gap = max(0.0, target - actual)

            new_rows.append(MachineDailyRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, unit=unit, section=section,
                shift=shift, machine_type=mtype, machine_id=mid, target_kg=target, actual_kg=actual,
                downtime_min=dt_min, downtime_reason=reason, is_planned_downtime=planned,
                stoppage_count=int(_num(rec.get("Stoppage Count"))), power_events=int(_num(rec.get("Power Events"))),
                energy_kwh=_opt_num(rec.get("Energy kWh")), quality_rating_pct=_opt_num(rec.get("Quality Rating %")),
                main_issue=_text(rec.get("Main Issue")), last_maintenance=_text(rec.get("Last Maintenance")),
                next_maintenance=_text(rec.get("Next Maintenance")),
            ))
            # Same row also feeds the Production page's existing tables
            new_rows.append(ProductionRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, shift=shift, department=section,
                unit=unit, machine_type=mtype, machine_id=mid, actual_kg=actual, target_kg=target, gap_kg=gap,
                achievement_pct=round(actual / target * 100, 1) if target > 0 else 0.0,
                source_sheet=MACHINE_SHEET, source_row=row_no, source_column="Actual Kg",
            ))
            if dt_min > 0 or reason:
                new_rows.append(DowntimeRecord(
                    id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, shift=shift,
                    machine_id=mid, machine_type=mtype, downtime_minutes=dt_min,
                    loss_kg=gap if reason else 0.0, reason_category=reason or "Machine Downtime",
                    remarks=_text(rec.get("Main Issue")),
                ))
            counts["machines"] = counts.get("machines", 0) + 1

        for row_no, rec in parsed.get(MANPOWER_SHEET, []):
            d = keep_date(MANPOWER_SHEET, row_no, rec)
            dept = _text(rec.get("Department"))
            if d is None:
                continue
            if not dept:
                warn(f"{MANPOWER_SHEET} row {row_no}: Department is required - row skipped.")
                continue
            req, avail = int(_num(rec.get("Required"))), int(_num(rec.get("Available")))
            new_rows.append(ManpowerRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d,
                shift=Normalizer.parse_shift(_text(rec.get("Shift"))), department=dept,
                planned_count=req, actual_count=avail, gap_count=max(0, req - avail),
            ))
            counts["manpower"] = counts.get("manpower", 0) + 1

        for row_no, rec in parsed.get(QUALITY_SHEET, []):
            d = keep_date(QUALITY_SHEET, row_no, rec)
            param = _text(rec.get("Parameter"))
            if d is None:
                continue
            if not param or _blank(rec.get("Value")) or _blank(rec.get("Limit")):
                warn(f"{QUALITY_SHEET} row {row_no}: Parameter, Value and Limit are required - row skipped.")
                continue
            new_rows.append(QualityParameterRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, machine_id=_text(rec.get("Machine ID")),
                parameter=param, value=_num(rec.get("Value")), limit_value=_num(rec.get("Limit")),
                unit=_text(rec.get("Unit")) or "",
            ))
            counts["quality"] = counts.get("quality", 0) + 1

        for row_no, rec in parsed.get(BUSINESS_SHEET, []):
            d = keep_date(BUSINESS_SHEET, row_no, rec)
            if d is None:
                continue
            new_rows.append(BusinessDailyRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d,
                revenue_lakhs=_num(rec.get("Revenue Lakhs")), orders_meters=_num(rec.get("Orders Meters")),
                dispatched_meters=_num(rec.get("Dispatched Meters")), collected_lakhs=_num(rec.get("Collected Lakhs")),
                outstanding_lakhs=_num(rec.get("Outstanding Lakhs")),
            ))
            counts["business"] = counts.get("business", 0) + 1

        for row_no, rec in parsed.get(STOCK_SHEET, []):
            d = keep_date(STOCK_SHEET, row_no, rec)
            cat = _text(rec.get("Category"))
            if d is None:
                continue
            if not cat:
                warn(f"{STOCK_SHEET} row {row_no}: Category is required - row skipped.")
                continue
            new_rows.append(StockRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, category=cat,
                current_value_lakhs=_num(rec.get("Current Value Lakhs")), limit_value_lakhs=_num(rec.get("Limit Value Lakhs")),
            ))
            counts["stock"] = counts.get("stock", 0) + 1

        for row_no, rec in parsed.get(ACTIONS_SHEET, []):
            d = keep_date(ACTIONS_SHEET, row_no, rec)
            issue, action = _text(rec.get("Issue")), _text(rec.get("Action"))
            if d is None:
                continue
            if not issue or not action:
                warn(f"{ACTIONS_SHEET} row {row_no}: Issue and Action are required - row skipped.")
                continue
            status = (_text(rec.get("Status")) or "OPEN").upper().replace(" ", "_")
            if status not in ("OPEN", "IN_PROGRESS", "COMPLETED"):
                warn(f"{ACTIONS_SHEET} row {row_no}: Status '{status}' not recognised - set to OPEN.")
                status = "OPEN"
            new_rows.append(ActionItemRecord(
                id=str(uuid.uuid4()), source_report_id=file_id, report_date=d, issue=issue, action=action,
                owner=_text(rec.get("Owner")) or "", status=status,
            ))
            counts["actions"] = counts.get("actions", 0) + 1

        if not counts:
            raise ValueError("The template contains no valid data rows. Fill at least one sheet and check the Date column.")

        # A new template replaces earlier template data for the same dates
        old_ids = [r[0] for r in db.query(ReportRecord.id).filter(ReportRecord.report_type == TEMPLATE_REPORT_TYPE).all()]
        if old_ids:
            for model in _REPLACEABLE_MODELS:
                db.query(model).filter(
                    model.source_report_id.in_(old_ids), model.report_date.in_(dates)
                ).delete(synchronize_session=False)

        db.add(ReportRecord(
            id=file_id, file_hash=file_hash, filename=filename, file_path=file_path,
            report_type=TEMPLATE_REPORT_TYPE, confidence_score=1.0, status="PROCESSED",
            report_date=max(dates), upload_timestamp=datetime.utcnow(),
        ))
        db.flush()
        db.add_all(new_rows)
        db.commit()

        summary = ", ".join(f"{n} {k}" for k, n in counts.items())
        return {
            "status": "SUCCESS",
            "message": f"Template imported: {summary}. All pages updated.",
            "report_id": file_id,
            "filename": filename,
            "detected_report_type": TEMPLATE_REPORT_TYPE,
            "final_report_type": TEMPLATE_REPORT_TYPE,
            "confidence_score": 1.0,
            "report_date": max(dates),
            "inserted_counts": counts,
            "validation_warnings": warnings,
        }
