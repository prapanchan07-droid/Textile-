"""Builds the downloadable Excel data template (empty, or pre-filled with sample data)."""
import io
import random
from datetime import date, timedelta
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.template.template_spec import (
    SHEETS, TEMPLATE_MARKER_SHEET, MACHINE_SHEET, MANPOWER_SHEET,
    QUALITY_SHEET, BUSINESS_SHEET, STOCK_SHEET, ACTIONS_SHEET,
)

REQUIRED_FILL = PatternFill("solid", fgColor="1E3A8A")
OPTIONAL_FILL = PatternFill("solid", fgColor="64748B")
HEADER_FONT = Font(bold=True, color="FFFFFF")

SAMPLE_END_DATE = date(2026, 9, 22)
SAMPLE_DAYS = 7

# machine_id: (type, section, unit, daily target kg, efficiency % per day, main reason, kWh per kg, quality %)
_MACHINES = {
    "V-09": ("Vortex", "Spinning Unit 1", "Unit I", 9100, [89.2, 87.5, 91.0, 82.1, 88.4, 86.0, 84.2], "Electrical / Power", 0.22, 96.1),
    "V-05": ("Vortex", "Spinning Unit 1", "Unit I", 9200, [91.5, 90.0, 89.2, 87.0, 88.5, 89.0, 88.1], "Maintenance", 0.21, 97.4),
    "V-12": ("Vortex", "Spinning Unit 1", "Unit I", 9200, [94.5, 95.0, 93.0, 94.2, 93.8, 94.0, 93.5], "Material / Process", 0.20, 98.6),
    "SMX-03": ("Simplex", "Prep Unit 2", "Unit II", 3390, [92.0, 91.8, 89.5, 90.0, 91.2, 89.8, 90.5], "Cleaning", 0.23, 97.8),
    "RF-04": ("Ring Frame", "Spinning Unit 2", "Unit II", 11500, [96.0, 95.5, 94.2, 95.0, 94.0, 95.2, 94.8], "Maintenance", 0.19, 99.1),
    "A-02": ("Airjet", "Weaving Unit 1", "Unit I", 10700, [96.2, 95.8, 96.0, 94.8, 95.5, 95.0, 95.2], "Other", 0.18, 99.3),
}
_QUALITY_SERIES = [  # (machine, parameter, limit, unit, 7-day values)
    ("V-09", "Thick/Km", 40, "/km", [38, 41, 44, 49, 48, 50, 52]),
    ("V-05", "U%", 12.0, "%", [11.2, 11.4, 11.5, 11.7, 11.6, 11.8, 11.8]),
    ("A-02", "Neps/Km", 140, "/km", [110, 115, 112, 125, 118, 122, 120]),
    ("RF-04", "Count CV", 1.5, "%", [1.1, 1.2, 1.1, 1.3, 1.2, 1.2, 1.2]),
    ("V-12", "Strength CV", 5.5, "%", [4.5, 4.6, 4.7, 4.9, 4.8, 4.7, 4.8]),
    ("V-12", "Thin/Km", 15, "/km", [7, 8, 8, 9, 8, 8, 8]),
]
_DEPARTMENTS = [("Spinning", 420, 390), ("Weaving", 300, 285), ("Maintenance", 80, 75),
                ("Sizing", 120, 118), ("Carding & Prep", 180, 178), ("Quality & Lab", 150, 149)]
_REVENUE = [3.9, 4.2, 4.5, 3.6, 4.1, 4.0, 4.2]


def _sample_rows() -> Dict[str, List[List[Any]]]:
    rng = random.Random(42)
    days = [SAMPLE_END_DATE - timedelta(days=SAMPLE_DAYS - 1 - i) for i in range(SAMPLE_DAYS)]
    rows: Dict[str, List[List[Any]]] = {name: [] for name in SHEETS}

    for d_idx, day in enumerate(days):
        for mid, (mtype, section, unit, daily_target, effs, reason, kwh_kg, qual) in _MACHINES.items():
            eff = effs[d_idx]
            for shift in ("Shift I", "Shift II", "Shift III"):
                target = round(daily_target / 3, 1)
                actual = round(target * (eff + rng.uniform(-1.5, 1.5)) / 100, 1)
                downtime = round(max(0.0, (100 - eff) * 5 + rng.uniform(-8, 8)), 0)
                power = 1 if reason == "Electrical / Power" and eff < 88 else 0
                rows[MACHINE_SHEET].append([
                    day, unit, section, shift, mtype, mid, target, actual, downtime,
                    reason if downtime > 0 else "", "Y" if reason == "Cleaning" else "N",
                    max(0, round(downtime / 30)), power, round(actual * kwh_kg, 0),
                    round(qual + rng.uniform(-0.5, 0.5), 1),
                    "Repeated motor trip" if mid == "V-09" and eff < 88 else "",
                    "2026-09-18", "2026-09-28",
                ])

        for dept, required, available in _DEPARTMENTS:
            rows[MANPOWER_SHEET].append([day, dept, "", required, available - (d_idx % 3) * (2 if required > 200 else 1)])
        rows[BUSINESS_SHEET].append([
            day, _REVENUE[d_idx], 30857, round(30857 * (0.62 + 0.03 * d_idx)), round(11.8 + (d_idx % 3) * 0.3, 1),
            round(84.0 + d_idx * 1.1, 1),
        ])
        for mid, param, limit, unit, series in _QUALITY_SERIES:
            rows[QUALITY_SHEET].append([day, mid, param, series[d_idx], limit, unit])

    last = days[-1]
    for cat, cur, lim in (("Fabric Stock", 12.34, 5.0), ("Yarn Stock - Unit I", 1.91, 3.0), ("Yarn Stock - Unit II", 2.15, 3.5)):
        rows[STOCK_SHEET].append([last, cat, cur, lim])
    rows[ACTIONS_SHEET].extend([
        [last + timedelta(days=1), "V-09 downtime", "Check motor alignment + power fluctuation events", "Maintenance Manager", "OPEN"],
        [last, "Power fluctuation", "Inspect transformer sub-station & main breaker", "Electrical Team", "IN_PROGRESS"],
        [last - timedelta(days=1), "Spinning shift 2 absenteeism", "Deploy pool contract workers to Spinning Unit 1", "HR & Manpower Head", "COMPLETED"],
    ])
    return rows


def _add_validation(ws, col_idx: int, formula: str):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(col_idx)
    dv.add(f"{letter}2:{letter}5000")


def _build_instructions(ws):
    ws.title = TEMPLATE_MARKER_SHEET
    ws["A1"] = "AI Factory Manager - Data Template"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [
        "1. Fill the sheets below (one row per record). Dark blue headers are required, grey headers are optional.",
        "2. Delete any rows you do not need. Do not rename sheets or column headers. Leave a sheet empty to skip it.",
        "3. Upload the file with 'Upload Reports'. Every page (Overview, Production, Machines, Comparison, Manpower & Quality, Revenue, Decision Center) updates from it.",
        "4. Uploading a new template replaces earlier template data for the same dates.",
        "5. Utilization = 100% - downtime / planned minutes (480 min per shift row, 1440 per full-day row).",
    ]
    for i, text in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=text)

    row = len(lines) + 5
    for name, spec in SHEETS.items():
        ws.cell(row=row, column=1, value=name).font = Font(bold=True, size=12)
        ws.cell(row=row, column=2, value=spec["purpose"])
        row += 1
        for c, title in enumerate(("Column", "Required", "Description", "Example"), start=1):
            cell = ws.cell(row=row, column=c, value=title)
            cell.font = HEADER_FONT
            cell.fill = OPTIONAL_FILL
        row += 1
        for col in spec["columns"]:
            ws.cell(row=row, column=1, value=col.header)
            ws.cell(row=row, column=2, value="Yes" if col.required else "No")
            ws.cell(row=row, column=3, value=col.description)
            ws.cell(row=row, column=4, value=col.example)
            row += 1
        row += 1
    for letter, width in (("A", 28), ("B", 12), ("C", 80), ("D", 22)):
        ws.column_dimensions[letter].width = width


def build_template_xlsx(with_sample: bool = False) -> bytes:
    wb = Workbook()
    _build_instructions(wb.active)
    sample = _sample_rows() if with_sample else {}

    for name, spec in SHEETS.items():
        ws = wb.create_sheet(name)
        for idx, col in enumerate(spec["columns"], start=1):
            cell = ws.cell(row=1, column=idx, value=col.header)
            cell.font = HEADER_FONT
            cell.fill = REQUIRED_FILL if col.required else OPTIONAL_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(idx)].width = max(14, len(col.header) + 4)
            if col.kind == "shift":
                _add_validation(ws, idx, '"Shift I,Shift II,Shift III"')
            elif col.kind == "yesno":
                _add_validation(ws, idx, '"Y,N"')
            elif col.kind == "status":
                _add_validation(ws, idx, '"OPEN,IN_PROGRESS,COMPLETED"')
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30

        for r_idx, row in enumerate(sample.get(name, []), start=2):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
        # Format the empty Date column so typed dates stay dates
        for r in range(len(sample.get(name, [])) + 2, 202):
            ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
