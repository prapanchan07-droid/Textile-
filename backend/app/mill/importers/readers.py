"""Read .xls (xlrd), .xlsx (openpyxl) and .csv into `Sheet` grids. Files are opened read-only and never modified."""
from __future__ import annotations

import csv
import io
from typing import List

from app.mill.importers.base import Sheet


class UnsupportedFile(ValueError):
    pass


def read_workbook(content: bytes, filename: str) -> List[Sheet]:
    name = filename.lower()
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # OLE2: legacy .xls whatever the extension says
        return _read_xls(content)
    if content[:2] == b"PK":
        return _read_xlsx(content)
    if name.endswith(".csv"):
        return _read_csv(content, filename)
    raise UnsupportedFile("Not an Excel (.xls/.xlsx) or CSV file")


def _read_xls(content: bytes) -> List[Sheet]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        row.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        out.append(Sheet(sh.name, rows))
    return out


def _read_xlsx(content: bytes) -> List[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [Sheet(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def _read_csv(content: bytes, filename: str) -> List[Sheet]:
    text = content.decode("utf-8-sig", errors="replace")
    rows = [list(r) for r in csv.reader(io.StringIO(text))]
    return [Sheet(filename.rsplit(".", 1)[0], [[v if v != "" else None for v in r] for r in rows])]

