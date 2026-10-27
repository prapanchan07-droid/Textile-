"""
GM-to-owner report: a snapshot of the period's numbers and work, plus the GM's commentary.

The snapshot is computed when the draft is created or refreshed and frozen at publication: later imports never
change a published report. Excel export renders the stored snapshot; PDF is the browser's print of the report page.
"""
from __future__ import annotations

import io
from datetime import date, datetime, time, timezone
from typing import Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.mill import constants as C
from app.mill.analytics import metrics as M
from app.mill.analytics.pages import command_center, energy, freshness, jsonable
from app.mill.analytics.records import Filters
from app.mill.models import Alert, ImportBatch, Task, WorkOrder

SECTIONS = ["production", "target", "efficiency", "deviations", "stops", "alerts", "maintenance", "groups", "tasks", "energy",
            "opportunities", "decisions", "risks"]


def build_snapshot(db: Session, start: date, end: date, shift: Optional[str] = None) -> dict:
    f = Filters(date_from=start, date_to=end, shifts=[shift] if shift else [])
    cc = command_center(db, f)
    k = cc["kpis"]
    t0 = datetime.combine(start, time.min, tzinfo=timezone.utc)
    t1 = datetime.combine(end, time.max, tzinfo=timezone.utc)
    alerts = list(db.scalars(select(Alert).where((Alert.status != "RESOLVED") | (Alert.resolved_at >= t0))))
    crit = [a for a in alerts if a.severity == "CRITICAL" or a.escalation_level > 0]
    tasks = list(db.scalars(select(Task)))
    now = datetime.now(timezone.utc)
    t_done = [t for t in tasks if t.completed_at and t0 <= t.completed_at <= t1]
    t_open = [t for t in tasks if t.status in C.TASK_OPEN]
    wos = list(db.scalars(select(WorkOrder).where((WorkOrder.reported_at <= t1) & ((WorkOrder.completed_at.is_(None)) | (WorkOrder.completed_at >= t0)))))
    batches = list(db.scalars(select(ImportBatch).where(ImportBatch.status == "COMMITTED", ImportBatch.period_end >= start, ImportBatch.period_start <= end)))
    findings = cc["findings"]
    stops = sorted((x for x in findings if x["kind"].startswith("HIGH_")), key=lambda x: x["title"])[:6]
    opportunities = [x for x in findings if x["basis"] == "MODELED"]
    en = energy(db, f)
    snap = {
        "meta": {"period": {"from": start.isoformat(), "to": end.isoformat(), "shift": shift}, "generated_at": now.isoformat(),
                 "freshness": freshness(db), "sources": cc["period"]["sources"], "days_with_data": cc["period"]["days_with_data"],
                 "import_batches": [{"id": b.id, "filename": b.filename, "format": b.source_format, "warnings": b.warning_count, "errors": b.error_count} for b in batches],
                 "definitions": {m.key: m.formula for m in M.DEFINITIONS}},
        "production": {"meters": k["meters"], "kilopicks": k["kilopicks"], "efficiency_pct": k["efficiency_pct"], "loom_shifts": k["loom_shifts"],
                       "reporting_looms": k["reporting_looms"], "registered_looms": k["registered_looms"],
                       "not_reporting_looms": k["not_reporting_looms"], "by_category": cc["by_category"]},
        "target": {"target_meters": k["target_meters"], "actual_meters": k["meters"],
                   "vs_target_pct": round((k["meters"] - k["target_meters"]) / k["target_meters"] * 100, 2) if k["target_meters"] and k["meters"] is not None else None,
                   "note": None if k["target_meters"] else "No production target configured for this period."},
        "efficiency": {"by_shift": {s: v["efficiency_pct"] for s, v in cc["by_shift"].items()},
                       "by_category": {c: v["efficiency_pct"] for c, v in cc["by_category"].items()}, "daily": cc["daily"]},
        "deviations": [x for x in findings if x["kind"] in ("BELOW_PEERS", "BELOW_OWN_BASELINE", "SHIFT_GAP")][:8],
        "stops": {"findings": stops, "totals": {"warp_stops": k["warp_stops"], "weft_stops": k["weft_stops"],
                                                "warp_stops_cmpx": k["warp_stops_cmpx"], "weft_stops_cmpx": k["weft_stops_cmpx"]},
                  "note": "Stop causes are known only from the Tsudakoma export's sensor counts; the mill report records counts, not reasons."},
        "alerts": {"open_by_severity": k["open_alerts"], "critical_or_escalated": [_alert(a) for a in crit[:15]],
                   "resolved_in_period": sum(1 for a in alerts if a.status == "RESOLVED")},
        "maintenance": {"open": [_wo(w) for w in wos if w.status in ("OPEN", "IN_PROGRESS", "ON_HOLD")][:20],
                        "completed_in_period": [_wo(w) for w in wos if w.completed_at and t0 <= w.completed_at <= t1][:20]},
        "groups": cc["groups"],
        "tasks": {"completed": len(t_done), "open": len(t_open), "blocked": sum(1 for t in t_open if t.status == "BLOCKED"),
                  "overdue": [_task(t) for t in t_open if t.due_at and t.due_at < now][:20],
                  "escalated_to_owner": [_task(t) for t in tasks if t.escalated_to_owner and t.status in C.TASK_OPEN]},
        "energy": {"air_status": en["air"]["status"], "energy_status": en["energy"]["status"], "requirements": en["requirements"]},
        "opportunities": [{"title": x["title"], "impact": x["impact"], "basis": x["basis"]} for x in opportunities][:8],
        "decisions": [_task(t) for t in tasks if t.escalated_to_owner and t.status in C.TASK_OPEN],
        "risks": ([f"{len(k['not_reporting_looms'])} registered loom(s) without records in the period"] if k["not_reporting_looms"] else [])
                 + [f"Import '{b.filename}': {b.warning_count} warning(s), {b.error_count} error(s)" for b in batches if b.warning_count or b.error_count]
                 + [f"Open critical alert: {a.title}" for a in crit if a.status != "RESOLVED"][:10],
    }
    return jsonable(snap)


def _alert(a: Alert) -> dict:
    return {"id": a.id, "title": a.title, "severity": a.severity, "status": a.status, "period": a.period_label, "escalation_level": a.escalation_level}


def _task(t: Task) -> dict:
    return {"id": t.id, "title": t.title, "status": t.status, "priority": t.priority, "due_at": t.due_at.isoformat() if t.due_at else None, "group_id": t.group_id}


def _wo(w: WorkOrder) -> dict:
    return {"id": w.id, "title": w.title, "kind": w.kind, "status": w.status, "loom_id": w.loom_id,
            "completed_at": w.completed_at.isoformat() if w.completed_at else None, "downtime_min": w.downtime_min}


def to_xlsx(report) -> bytes:
    """Excel copy of a report's stored snapshot (never recomputed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    s = report.snapshot or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    p = s.get("meta", {}).get("period", {})
    lines = [
        ("Report", report.title), ("Status", report.status), ("Period", f"{p.get('from')} to {p.get('to')}" + (f" shift {p.get('shift')}" if p.get("shift") else "")),
        ("Generated", s.get("meta", {}).get("generated_at")), ("Published", report.published_at.isoformat() if report.published_at else "not published"),
        ("", ""), ("Meters", s.get("production", {}).get("meters")), ("Kilopicks", s.get("production", {}).get("kilopicks")),
        ("Efficiency % (time-weighted)", s.get("production", {}).get("efficiency_pct")),
        ("Target meters", s.get("target", {}).get("target_meters")), ("vs target %", s.get("target", {}).get("vs_target_pct")),
        ("", ""), ("GM commentary", report.commentary or ""), ("Decisions needed", report.decisions_needed or ""), ("Risks", report.risks or ""),
    ]
    for a, b in lines:
        ws.append([a, b])
        ws.cell(ws.max_row, 1).font = bold
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 90

    def table(title, header, rows):
        sh = wb.create_sheet(title)
        sh.append(header)
        for c in sh[1]:
            c.font = bold
        for r in rows:
            sh.append(r)

    table("By category", ["Category", "Looms", "Loom-shifts", "Meters", "Kilopicks", "Efficiency %", "Warp /100k picks", "Weft /100k picks"],
          [[c, v.get("looms"), v.get("loom_shifts"), v.get("meters"), v.get("kilopicks"), v.get("efficiency_pct"), v.get("warp_stops_cmpx"), v.get("weft_stops_cmpx")]
           for c, v in s.get("production", {}).get("by_category", {}).items()])
    table("Daily", ["Date", "Meters", "Efficiency %"], [[d["date"], d.get("meters"), d.get("efficiency_pct")] for d in s.get("efficiency", {}).get("daily", [])])
    table("Deviations", ["Finding", "Basis", "Next step"], [[x["title"], x["basis"], x.get("next_step")] for x in s.get("deviations", [])])
    table("Groups", ["Group", "Leader", "Looms", "Meters", "Efficiency %", "Open tasks", "Blocked", "Open alerts"],
          [[g["name"], g.get("leader"), g.get("looms"), g.get("meters"), g.get("efficiency_pct"), g.get("open_tasks"), g.get("blocked_tasks"), g.get("open_alerts")]
           for g in s.get("groups", [])])
    table("Alerts", ["Alert", "Severity", "Status", "Period", "Escalation"],
          [[a["title"], a["severity"], a["status"], a.get("period"), a.get("escalation_level")] for a in s.get("alerts", {}).get("critical_or_escalated", [])])
    table("Risks", ["Risk"], [[r] for r in s.get("risks", [])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
